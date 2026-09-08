from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'docs/Entrega_Ciclo_de_Solicitud.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Inches(.7), Inches(.7)
sec.left_margin, sec.right_margin = Inches(.8), Inches(.8)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10)
for name in ['Title','Heading 1','Heading 2']:
    styles[name].font.name = 'Aptos'; styles[name].font.color.rgb = RGBColor(0,0,0)
styles['Title'].font.size = Pt(22)

def title(t):
    p=doc.add_paragraph(style='Title'); p.add_run(t); return p
def h(t, level=1): doc.add_heading(t, level=level)
def p(t=''):
    x=doc.add_paragraph(t); x.paragraph_format.space_after=Pt(6); return x
def code(t):
    x=doc.add_paragraph(); r=x.add_run(t); r.font.name='Consolas'; r.font.size=Pt(9); return x
def shade(cell, fill='1F4E78'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def table(headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c, text in zip(t.rows[0].cells, headers):
        c.text=text; shade(c)
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.font.bold=True
    for row in rows:
        cells=t.add_row().cells
        for c, text in zip(cells,row): c.text=str(text); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()

title('Entrega Ciclo de Solicitud')
p('Documentación técnica de una API de gestión de incidencias. El proyecto preserva datos al evolucionar su base PostgreSQL mediante migraciones versionadas.')
h('Objetivo')
p('Construir una API REST para crear y consultar incidencias y categorías, con validación, documentación Swagger, persistencia con Prisma y PostgreSQL. La entrega demuestra cómo evolucionar una relación obligatoria sin perder los datos que existían antes del cambio.')
h('Tecnologías')
table(['Capa','Tecnología'], [('Lenguaje','TypeScript'),('API','NestJS'),('Persistencia','Prisma ORM y PostgreSQL 16'),('Contenedor','Docker Compose'),('Contrato HTTP','Swagger OpenAPI'),('Validación','class-validator')])
h('Arquitectura y recorrido de una solicitud')
p('El controller no concentra reglas de negocio: recibe HTTP y delega. El service verifica que la categoría exista; el repository contiene las consultas Prisma.')
code('POST /incidents  ->  DTO + ValidationPipe  ->  Controller  ->  Service')
code('                 ->  Repository  ->  Prisma Client  ->  PostgreSQL  ->  JSON HTTP')
p('El DTO controla los campos, tipos y estados permitidos. Los errores de validación son 400; recursos inexistentes son 404; una violación de unicidad de Prisma se transforma en 409; los demás errores Prisma se responden como 500.')
h('Modelo de datos')
table(['Entidad','Campos'], [('Category','id, code UNIQUE, name, createdAt'),('Incident','id, title, description, status, categoryId FK, createdAt, updatedAt')])
code('Category 1  ----------------  N Incident')
p('categoryId es obligatorio al finalizar la evolución y tiene una clave foránea hacia Category.id.')
h('Migraciones y conservación de datos')
table(['Orden','Migración','Propósito'], [('1','20260908154500_init_incident','Crea Incident e inserta dos incidencias históricas.'),('2','20260908160000_add_category_and_optional_relation','EXPANDIR: crea Category, UNIQUE code, FK y categoryId opcional.'),('3','20260908161000_move_legacy_incidents_to_general','MIGRAR DATOS: crea GENERAL y actualiza las incidencias antiguas.'),('4','20260908162000_make_incident_category_required','CONTRAER: hace obligatorio categoryId.')])
p('La secuencia expandir -> migrar datos -> contraer evita una alteración incompatible sobre filas existentes. Las migraciones anteriores permanecen intactas después de aplicarse.')
h('Seed y restricciones')
p('El seed usa upsert para GENERAL y SOFTWARE. La incidencia de ejemplo se busca y actualiza si existe, por lo que npm run seed se puede repetir sin duplicar información.')
p('PostgreSQL impone Category.code UNIQUE. Crear dos categorías con el mismo code produce el error P2002 de Prisma, que el filtro de excepciones expone como HTTP 409.')
h('Endpoints y Swagger')
table(['Método','Ruta','Acción'], [('POST','/incidents','Crea una incidencia'),('GET','/incidents','Lista incidencias con categoría'),('GET','/incidents/:id','Consulta una incidencia'),('POST','/categories','Crea una categoría'),('GET','/categories','Lista categorías')])
p('La UI de Swagger se publica en /api cuando la aplicación está en ejecución.')
h('Pruebas y reconstrucción')
p('Evidencia obtenida en esta entrega: npm run prisma:generate completó correctamente y npm run build completó correctamente después de crear los módulos, DTOs, filtro y repositorios.')
p('La reconstrucción prevista desde cero es: docker compose down -v; docker compose up -d; npx prisma migrate deploy; npm run seed. Durante la entrega el proceso de Docker Desktop quedó iniciado, pero su API Linux no estuvo disponible; por integridad, no se registra como prueba ejecutada. El archivo de evidencia se actualizará al ejecutar esos comandos con el motor disponible.')
h('Desarrollo y despliegue')
p('prisma migrate dev es para desarrollo: genera migraciones desde cambios del esquema. prisma migrate deploy no crea ni modifica migraciones; únicamente aplica las versionadas. Por esa razón se usa para despliegue y reconstrucción.')
h('Repositorio')
p('https://github.com/jefroMMM/Ciclo-de-solicitud')
doc.save(OUT)
